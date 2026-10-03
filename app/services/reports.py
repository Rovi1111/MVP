"""Единый отчёт по сохранённому анализу: экспорт в Excel и PDF.

PDF строится ReportLab: он ставится через pip без системных библиотек
(WeasyPrint потребовал бы pango/cairo из Homebrew).
"""

import io
import os
import re
from functools import lru_cache
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy.orm import Session

from app.errors import AppError
from app.models import Analysis, User
from app.services import analyses, charts, mailer
from app.services.ai_summary import money, percent
from app.services.audit import log_action
from app.services.settings import get_org_settings
from app.timeutil import to_local

CHANGE_LABELS = {
    "increase": "Рост",
    "decrease": "Снижение",
    "new": "Новая",
    "removed": "Исчезнувшая",
    "no_change": "Без изменений",
}
MONEY_FORMAT = "#,##0.00"
PCT_FORMAT = '0.0"%"'
_HEADER_FILL = PatternFill("solid", fgColor="DDE5F4")
_ALERT_FILL = PatternFill("solid", fgColor="FDECEA")

# Шрифт с кириллицей для PDF: (обычный, полужирный). Первый найденный вариант.
_STATIC_FONTS = Path(__file__).resolve().parent.parent / "static" / "fonts"
_FONT_CANDIDATES = [
    (_STATIC_FONTS / "DejaVuSans.ttf", _STATIC_FONTS / "DejaVuSans-Bold.ttf"),
    (
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ),
    (Path("/Library/Fonts/Arial.ttf"), Path("/Library/Fonts/Arial Bold.ttf")),
    (
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ),
    (
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ),
]


def file_name(record: Analysis, extension: str) -> str:
    return f"report_{record.id}_{to_local(record.created_at):%Y%m%d}.{extension}"


def _param_rows(record: Analysis, organization: str) -> list[tuple[str, str]]:
    snapshot = record.result
    rows = [
        ("Организация", organization),
        ("Отчёт", record.title),
        ("Дата расчёта", f"{to_local(record.created_at):%d.%m.%Y %H:%M}"),
        ("Автор", record.author.login),
    ]
    comparison = snapshot.get("comparison")
    if comparison:
        params = comparison["params"]
        for title, key in (("Период 1", "period1"), ("Период 2", "period2")):
            period = params[key]
            rows.append((title, f"{period['label']} ({period['start']} — {period['end']})"))
        rows.append(("ЦФО", params.get("cost_center") or "Все"))
        if params.get("category") or params.get("change_types"):
            types = ", ".join(CHANGE_LABELS[t] for t in params.get("change_types") or [])
            rows.append(
                ("Фильтры таблицы", "; ".join(x for x in (params.get("category"), types) if x))
            )
    chart = snapshot.get("chart")
    if chart:
        version = chart.get("budget_version")
        rows.append(("Диаграмма", f"факт и бюджет, {chart['period']['label']}"))
        rows.append(("Версия бюджета", str(version) if version else "бюджет не загружен"))
    return rows


# ---------- Excel ----------


def _sheet(workbook: Workbook, title: str, headers: list[str], widths: list[int]):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"
    return sheet


def _format_columns(sheet, formats: dict[int, str]) -> None:
    for row in sheet.iter_rows(min_row=2):
        for index, number_format in formats.items():
            if isinstance(row[index].value, int | float):
                row[index].number_format = number_format


def build_excel(record: Analysis, organization: str) -> bytes:
    """Книга Excel: параметры, сравнение, данные диаграммы, AI-заключение. Суммы — числами."""
    snapshot = record.result
    workbook = Workbook()
    workbook.remove(workbook.active)

    sheet = _sheet(workbook, "Параметры", ["Параметр", "Значение"], [28, 70])
    for row in _param_rows(record, organization):
        sheet.append(row)

    comparison = snapshot.get("comparison")
    if comparison:
        params = comparison["params"]
        headers = [
            "Категория", params["period1"]["label"], params["period2"]["label"],
            "Изменение", "Изменение, %", "Тип изменения",
        ]  # fmt: skip
        sheet = _sheet(workbook, "Сравнение", headers, [36, 18, 18, 18, 14, 18])
        for row in comparison["rows"]:
            sheet.append(
                [
                    row["category"], row["amount_p1"], row["amount_p2"], row["delta"],
                    row["pct_change"] if row["pct_change"] is not None else "н/д",
                    CHANGE_LABELS[row["change_type"]],
                ]
            )  # fmt: skip
        sheet.append(
            [
                "Итого", comparison["total_p1"], comparison["total_p2"], comparison["total_delta"],
                comparison["total_pct"] if comparison["total_pct"] is not None else "н/д", "",
            ]
        )  # fmt: skip
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        _format_columns(sheet, {1: MONEY_FORMAT, 2: MONEY_FORMAT, 3: MONEY_FORMAT, 4: PCT_FORMAT})

    chart = snapshot.get("chart")
    if chart:
        headers = (
            ["Месяц"] + [item["category"] for item in chart["series"]] + ["Факт всего", "Бюджет"]
        )
        sheet = _sheet(workbook, "Факт и бюджет", headers, [18] + [18] * (len(headers) - 1))
        for index, label in enumerate(chart["periods"]):
            budget = chart["budget"][index] if chart.get("budget") is not None else "нет"
            amounts = [item["amounts"][index] for item in chart["series"]]
            sheet.append([label, *amounts, chart["totals"][index], budget])
        _format_columns(sheet, dict.fromkeys(range(1, len(headers)), MONEY_FORMAT))

    sheet = _sheet(workbook, "AI-заключение", ["AI-заключение"], [120])
    sheet.append([record.ai_text or "AI-анализ не запускался"])
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------- PDF ----------


@lru_cache
def _fonts() -> tuple[str, str]:
    """Регистрирует шрифт с кириллицей и возвращает имена (обычный, полужирный)."""
    candidates = list(_FONT_CANDIDATES)
    if os.environ.get("PDF_FONT_PATH"):
        regular = Path(os.environ["PDF_FONT_PATH"])
        candidates.insert(0, (regular, Path(os.environ.get("PDF_FONT_BOLD_PATH", regular))))
    for regular, bold in candidates:
        if regular.exists():
            pdfmetrics.registerFont(TTFont("Report", str(regular)))
            pdfmetrics.registerFont(TTFont("Report-Bold", str(bold if bold.exists() else regular)))
            return "Report", "Report-Bold"
    raise AppError(
        "Не найден шрифт с кириллицей для PDF. Положите DejaVuSans.ttf в app/static/fonts/ "
        "или укажите путь в переменной PDF_FONT_PATH"
    )


def _styles() -> dict[str, ParagraphStyle]:
    regular, bold = _fonts()
    base = ParagraphStyle("base", fontName=regular, fontSize=9, leading=12)
    return {
        "base": base,
        "num": ParagraphStyle("num", parent=base, alignment=TA_RIGHT),
        "bold": ParagraphStyle("bold", parent=base, fontName=bold),
        "num_bold": ParagraphStyle("num_bold", parent=base, fontName=bold, alignment=TA_RIGHT),
        "title": ParagraphStyle("title", parent=base, fontName=bold, fontSize=16, leading=20),
        "h2": ParagraphStyle(
            "h2", parent=base, fontName=bold, fontSize=12, leading=16, spaceBefore=10
        ),
        "text": ParagraphStyle("text", parent=base, fontSize=10, leading=14, spaceAfter=4),
    }


def _table(rows: list[list], widths: list[float], highlight: set[int] = frozenset()) -> Table:
    table = Table(rows, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE5F4")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C5CBD5")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEABOVE", (0, -1), (-1, -1), 1, colors.HexColor("#6B7482")),
    ]
    style += [("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FDECEA")) for i in highlight]
    table.setStyle(TableStyle(style))
    return table


def _image(figure, width_mm: float, pixels: tuple[int, int]) -> Image:
    png = charts.figure_png(figure, width=pixels[0], height=pixels[1], scale=2)
    return Image(io.BytesIO(png), width=width_mm * mm, height=width_mm * mm * pixels[1] / pixels[0])


def _ai_paragraphs(text: str, styles: dict) -> list:
    """Текст модели → абзацы; **жирный** и заголовки Markdown отображаются как оформление."""
    flow = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        heading = line.startswith("#")
        line = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escape(line.lstrip("# ")))
        flow.append(Paragraph(line, styles["bold"] if heading else styles["text"]))
    return flow


def build_pdf(record: Analysis, organization: str) -> bytes:
    """PDF-отчёт: параметры, таблица сравнения, диаграмма-картинка, AI-заключение."""
    styles = _styles()
    snapshot = record.result
    figures = analyses.snapshot_figures(snapshot)
    width = A4[0] - 30 * mm

    def p(text, style="base") -> Paragraph:
        return Paragraph(escape(str(text)), styles[style])

    flow: list = [p(record.title, "title"), Spacer(1, 4 * mm)]
    params = [[p(k, "bold"), p(v)] for k, v in _param_rows(record, organization)]
    flow.append(Table(params, colWidths=[45 * mm, width - 45 * mm]))

    comparison = snapshot.get("comparison")
    if comparison:
        labels = comparison["params"]
        flow.append(p("Сравнение периодов", "h2"))
        rows = [
            [
                p("Категория", "bold"), p(labels["period1"]["label"], "num_bold"),
                p(labels["period2"]["label"], "num_bold"), p("Изменение", "num_bold"),
                p("%", "num_bold"), p("Тип", "bold"),
            ]
        ]  # fmt: skip
        for row in comparison["rows"]:
            rows.append(
                [
                    p(row["category"]), p(money(row["amount_p1"]), "num"),
                    p(money(row["amount_p2"]), "num"), p(money(row["delta"]), "num"),
                    p(percent(row["pct_change"]), "num"), p(CHANGE_LABELS[row["change_type"]]),
                ]
            )  # fmt: skip
        rows.append(
            [
                p("Итого", "bold"), p(money(comparison["total_p1"]), "num_bold"),
                p(money(comparison["total_p2"]), "num_bold"),
                p(money(comparison["total_delta"]), "num_bold"),
                p(percent(comparison["total_pct"]), "num_bold"), p(""),
            ]
        )  # fmt: skip
        flow.append(_table(rows, [width * s for s in (0.26, 0.155, 0.155, 0.155, 0.105, 0.17)]))
        for warning in comparison.get("warnings", []):
            flow.append(p(f"Внимание: {warning}"))
        flow.append(Spacer(1, 4 * mm))
    if "fact_budget" in figures:
        flow.append(p("Факт и бюджет по месяцам", "h2"))
        flow.append(_image(figures["fact_budget"], 180, (1100, 560)))

    flow.append(p("AI-заключение", "h2"))
    if record.ai_text:
        flow += _ai_paragraphs(record.ai_text, styles)
    else:
        flow.append(p("AI-анализ не запускался.", "text"))

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer, pagesize=A4, title=record.title, author=organization,
        leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
    ).build(flow)  # fmt: skip
    return buffer.getvalue()


FORMATS = {
    "pdf": (build_pdf, "application/pdf"),
    "xlsx": (build_excel, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
}


def export(record: Analysis, organization: str, fmt: str) -> tuple[bytes, str, str]:
    """Возвращает (содержимое, MIME-тип, имя файла)."""
    try:
        builder, mime = FORMATS[fmt]
    except KeyError:
        raise AppError("Формат отчёта: pdf или xlsx") from None
    return builder(record, organization), mime, file_name(record, fmt)


def export_report(db: Session, user: User, analysis_id: int, fmt: str) -> tuple[bytes, str, str]:
    """Экспорт сохранённого анализа с записью в журнал: (содержимое, MIME-тип, имя файла)."""
    record = analyses.get_analysis(db, user, analysis_id)
    result = export(record, user.organization.name, fmt)
    log_action(db, user, "export", "analysis", record.id, {"format": fmt})
    db.commit()
    return result


def email_report(db: Session, user: User, analysis_id: int, to: str, formats: list[str]) -> None:
    """Отправляет отчёт на указанный адрес вложениями в выбранных форматах."""
    record = analyses.get_analysis(db, user, analysis_id)
    settings = get_org_settings(db, user.org_id)
    organization = user.organization.name
    attachments = [export(record, organization, fmt) for fmt in formats or ["pdf"]]
    body = f"{organization}\n{record.title}\n\nОтчёт во вложении. Отправил: {user.login}."
    if record.ai_text:
        body += f"\n\nAI-заключение:\n{record.ai_text}"
    mailer.send_mail(settings, to, f"Отчёт о расходах: {record.title}", body, attachments)
    log_action(db, user, "email", "analysis", record.id, {"to": to.strip(), "formats": formats})
    db.commit()
