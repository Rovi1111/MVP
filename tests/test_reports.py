"""Отчёты: экспорт PDF/Excel, отправка на email, журнал аудита, управление пользователями."""

import io
import smtplib

import pytest
from openpyxl import load_workbook

from app.services import mailer
from tests.helpers import load

FACT = [
    ("15.01.2025", "Аренда", 100000, "Офис"),
    ("2025-01", "Реклама", 40000.55, "Офис"),
    ("10.02.2025", "Аренда", 180000, "Офис"),
    ("11.02.2025", "Лицензии", 25000, "Склад"),
]
BUDGET = [("2025-02", "Аренда", 150000), ("2025-02", "Лицензии", 30000)]


@pytest.fixture
def client(client_for):
    client = client_for("a_admin")
    load(client, "fact", FACT)
    load(client, "budget", BUDGET)
    return client


@pytest.fixture
def analysis_id(client, db):
    from app.models import Analysis

    created = client.post(
        "/api/analyses", json={"kind": "comparison", "period1": "2025-01", "period2": "2025-02"}
    ).json()
    record = db.get(Analysis, created["id"])
    record.ai_text = "# Вывод\nРасходы **выросли**.\n\nРекомендации: пересмотреть аренду <офис>."
    db.commit()
    return created["id"]


class FakeSMTP:
    sent: list = []
    fail = False

    def __init__(self, host, port, timeout=None):
        self.host, self.port, self.steps = host, port, []

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def starttls(self):
        self.steps.append("tls")

    def login(self, user, password):
        self.steps.append(("login", user, password))

    def send_message(self, message):
        if FakeSMTP.fail:
            raise smtplib.SMTPRecipientsRefused({})
        FakeSMTP.sent.append((self, message))


@pytest.fixture
def smtp(monkeypatch):
    FakeSMTP.sent, FakeSMTP.fail = [], False
    monkeypatch.setattr(mailer.smtplib, "SMTP", FakeSMTP)
    return FakeSMTP


SMTP_SETTINGS = {
    "smtp_host": "smtp.example.ru", "smtp_port": 2525, "smtp_user": "robot",
    "smtp_password": "smtp-secret", "smtp_from": "reports@example.ru",
}  # fmt: skip


def test_excel_export_has_numeric_cells(client, analysis_id):
    response = client.get(f"/api/analyses/{analysis_id}/export", params={"format": "xlsx"})
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert f"report_{analysis_id}_" in response.headers["content-disposition"]

    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Параметры", "Сравнение", "Факт и бюджет", "AI-заключение"]
    rows = {row[0]: row for row in workbook["Сравнение"].iter_rows(min_row=2, values_only=True)}
    assert rows["Аренда"][1:5] == (100000, 180000, 80000, 80)
    assert rows["Реклама"][1] == 40000.55 and isinstance(
        rows["Реклама"][1], float
    )  # число, не текст
    assert rows["Лицензии"][4] == "н/д" and rows["Лицензии"][5] == "Новая"
    assert rows["Итого"][1:4] == (140000.55, 205000, 64999.45)
    assert workbook["Сравнение"]["B2"].number_format == "#,##0.00"

    chart = list(workbook["Факт и бюджет"].iter_rows(values_only=True))
    assert chart[0] == ("Месяц", "Аренда", "Реклама", "Лицензии", "Факт всего", "Бюджет")
    assert chart[1] == ("январь 2025", 100000, 40000.55, 0, 140000.55, 0)
    assert chart[2] == ("февраль 2025", 180000, 0, 25000, 205000, 180000)
    assert "выросли" in workbook["AI-заключение"]["A2"].value
    params = dict(workbook["Параметры"].iter_rows(min_row=2, values_only=True))
    assert params["Организация"] == "Организация А" and params["Версия бюджета"] == "1"


def test_pdf_export_with_single_chart(client, analysis_id):
    response = client.get(f"/api/analyses/{analysis_id}/export", params={"format": "pdf"})
    assert response.status_code == 200 and response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    # одна диаграмма «факт и бюджет» (PNG с прозрачностью хранится как картинка + маска)
    assert 1 <= response.content.count(b"/Subtype /Image") <= 2
    assert len(response.content) > 30_000


def test_export_without_ai_and_budget(client_for):
    client = client_for("b_admin")
    load(client, "fact", [("2025-01", "Аренда", 10), ("2025-02", "Аренда", 20)])
    created = client.post("/api/analyses", json={"period1": "2025-01", "period2": "2025-02"}).json()
    assert created["result"]["plan_fact"] is None and created["result"]["chart"]["budget"] is None
    assert client.get(f"/api/analyses/{created['id']}/export").content.startswith(b"%PDF")
    workbook = load_workbook(
        io.BytesIO(client.get(f"/api/analyses/{created['id']}/export?format=xlsx").content)
    )
    assert workbook["AI-заключение"]["A2"].value == "AI-анализ не запускался"
    assert [c.value for c in workbook["Факт и бюджет"][2]] == ["январь 2025", 10, 10, "нет"]
    params = dict(workbook["Параметры"].iter_rows(min_row=2, values_only=True))
    assert params["Версия бюджета"] == "бюджет не загружен"


def test_export_errors_and_isolation(client, client_for, analysis_id):
    assert client.get(f"/api/analyses/{analysis_id}/export?format=docx").status_code == 422
    assert client.get("/api/analyses/999/export").status_code == 404
    assert client_for("b_admin").get(f"/api/analyses/{analysis_id}/export").status_code == 404
    assert (
        client_for("a_employee").get(f"/api/analyses/{analysis_id}/export?format=xlsx").status_code
        == 200
    )


def test_email_requires_smtp_settings(client, analysis_id, smtp):
    response = client.post(f"/api/analyses/{analysis_id}/email", json={"to": "boss@example.ru"})
    assert response.status_code == 422 and "не настроена" in response.json()["detail"]
    assert smtp.sent == []
    page = client.get(f"/web/analyses/{analysis_id}")
    assert "SMTP не настроен" in page.text and "Отправка недоступна" in page.text


def test_email_report(client, analysis_id, smtp, db):
    client.patch("/api/settings", json=SMTP_SETTINGS)
    assert (
        client.post(f"/api/analyses/{analysis_id}/email", json={"to": "не адрес"}).status_code
        == 422
    )

    response = client.post(
        f"/api/analyses/{analysis_id}/email",
        json={"to": " boss@example.ru ", "formats": ["pdf", "xlsx"]},
    )
    assert (
        response.status_code == 200
        and response.json()["detail"] == "Отчёт отправлен на boss@example.ru"
    )
    server, message = smtp.sent[0]
    assert (server.host, server.port) == ("smtp.example.ru", 2525)
    assert server.steps == ["tls", ("login", "robot", "smtp-secret")]
    assert message["To"] == "boss@example.ru" and message["From"] == "reports@example.ru"
    assert "январь 2025" in message["Subject"]
    names = [part.get_filename() for part in message.iter_attachments()]
    assert [name.rsplit(".", 1)[1] for name in names] == ["pdf", "xlsx"]

    web = client.post(
        f"/web/analyses/{analysis_id}/email", data={"to": "cfo@example.ru", "formats": "pdf"}
    )
    assert "Отчёт отправлен на cfo@example.ru" in web.text and len(smtp.sent) == 2
    assert "Отправка недоступна" not in client.get(f"/web/analyses/{analysis_id}").text

    smtp.fail = True
    failed = client.post(f"/api/analyses/{analysis_id}/email", json={"to": "boss@example.ru"})
    assert failed.status_code == 502 and "Не удалось отправить письмо" in failed.json()["detail"]


def test_audit_log_records_actions(client, analysis_id, client_for, smtp):
    client.patch("/api/settings", json=SMTP_SETTINGS)
    client.get(f"/api/analyses/{analysis_id}/export?format=xlsx")
    client.post(f"/api/analyses/{analysis_id}/email", json={"to": "boss@example.ru"})
    client.post("/api/categories", json={"name": "Прочее"})
    client.post("/api/auth/logout")
    client.login("a_admin")

    entries = client.get("/api/audit").json()
    seen = {(e["action"], e["entity"]) for e in entries}
    for expected in [
        ("login", "session"),
        ("logout", "session"),
        ("create", "upload"),
        ("create", "category"),
        ("create", "analysis"),
        ("export", "analysis"),
        ("email", "analysis"),
        ("update", "settings"),
    ]:
        assert expected in seen, expected
    assert entries[0]["action"] == "login" and entries[0]["user_login"] == "a_admin"  # новые сверху
    settings_entry = next(e for e in entries if e["entity"] == "settings")
    assert settings_entry["details"]["smtp_password"] == "***"  # секреты в журнал не попадают
    assert all(
        e["action"] == "export"
        for e in client.get("/api/audit", params={"action": "export"}).json()
    )

    page = client.get("/web/audit")
    assert (
        page.status_code == 200
        and "Отправка на email" in page.text
        and "boss@example.ru" in page.text
    )
    assert "Экспорт" in client.get("/web/audit", params={"action": "export"}).text

    assert client_for("a_employee").get("/api/audit").status_code == 403
    assert client_for("a_employee").get("/web/audit").status_code == 403
    other = client_for("b_admin").get("/api/audit").json()
    assert {e["user_login"] for e in other} == {"b_admin"}


def test_users_page(client, client_for):
    page = client.get("/web/users")
    assert page.status_code == 200 and "a_employee" in page.text and "b_admin" not in page.text

    created = client.post(
        "/web/users",
        data={
            "login": "newbie",
            "full_name": "Новый",
            "password": "long-password",
            "role": "employee",
        },
    )
    assert "newbie" in created.text
    newbie = next(u for u in client.get("/api/users").json() if u["login"] == "newbie")
    assert client_for().login("newbie", "long-password").status_code == 200

    updated = client.post(
        f"/web/users/{newbie['id']}",
        data={"full_name": "Новый", "role": "manager", "password": "another-password"},
    )  # флажок «Активен» снят — пользователь отключается
    assert updated.status_code == 200
    after = next(u for u in client.get("/api/users").json() if u["login"] == "newbie")
    assert (after["role"], after["is_active"]) == ("manager", False)
    assert client_for().login("newbie", "another-password").status_code == 401

    short = client.post(
        "/web/users",
        data={"login": "x", "password": "123", "role": "employee"},
        headers={"HX-Request": "true"},
    )
    assert short.status_code == 422 and "не короче" in short.text
    assert client_for("a_employee").get("/web/users").status_code == 403


def test_menu_matches_permissions(client_for):
    admin_menu = client_for("a_admin").get("/web/dashboard").text
    employee_menu = client_for("a_employee").get("/web/dashboard").text
    for title in ("Анализ расходов", "Сохранённые анализы", "Ввод данных", "Загрузки"):
        assert title in admin_menu and title in employee_menu
    for title in ("Справочники", "Настройки", "Пользователи", "Журнал"):
        assert title in admin_menu and title not in employee_menu
